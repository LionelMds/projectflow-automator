from __future__ import annotations

import stat
import sys
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from projectflow.cad import templates
from projectflow.cad.demo_document_manager import (
    DEMO_LICENSE_KEY,
    JsonDocument,
    JsonDocumentManager,
    open_json_document_manager,
    read_demo_document,
    write_demo_document,
)
from projectflow.cad.document_manager import open_document_manager
from projectflow.cad.license_storage import SolidWorksLicenseStorage
from projectflow.cad.solidworks_properties import compute_property_updates
from projectflow.cad.templates import (
    CadTemplateService,
    cad_destination_dir,
    copy_new_file,
    plan_template_copy,
    template_availability,
)
from projectflow.config import AppConfig, CadConfig, CadPropertyNames
from projectflow.core.fiche_service import FicheService
from projectflow.core.models import CadOutcome, ProjectInput
from projectflow.core.numero import parse_project_number
from projectflow.core.project_service import ProjectService
from projectflow.exceptions import CadUnavailableError

TEMPLATE_PARTS = ["ENV-100.SLDPRT", "PRT-100.SLDPRT", "PRT-200.SLDPRT"]
TEMPLATE_PROPERTIES = {
    "Projet": "20XX-XXXX",
    "Client": "",
    "Auteur": "",
    "Description": "",
    "Révision": "",
    "Fournisseur": "Balz Métal SA",
    "FaireouAcheter": "Faire",
    "Finition": "Galvanisé",
    "TAG": "Repère 20XX-XXXX",
}


def _solidworks_templates(root: Path, *, reference_root: Path | None = None) -> Path:
    template_dir = root / "11-Racine Solidworks"
    stored_root = reference_root or template_dir
    for name in TEMPLATE_PARTS:
        write_demo_document(template_dir / f"20XX-XXXX-{name}", TEMPLATE_PROPERTIES)
    write_demo_document(
        template_dir / "20XX-XXXX-ENS-100.SLDASM",
        TEMPLATE_PROPERTIES,
        [f"{stored_root}\\20XX-XXXX-{name}" for name in TEMPLATE_PARTS]
        if reference_root is not None
        else [str(template_dir / f"20XX-XXXX-{name}") for name in TEMPLATE_PARTS],
    )
    return template_dir


def _autocad_templates(root: Path) -> Path:
    template_dir = root / "12-Racine AutoCAD"
    template_dir.mkdir(parents=True)
    (template_dir / "20XX-XXXX-ENS-100.dwg").write_bytes(b"dwg")
    (template_dir / "20XX-XXXX-ENS-100.dwl").write_bytes(b"lock")
    (template_dir / "20XX-XXXX-ENS-100.dwl2").write_bytes(b"lock")
    (template_dir / "20XX-XXXX-ENS-100.bak").write_bytes(b"backup")
    return template_dir


def _config(tmp_path: Path, *, subfolder: str = "") -> CadConfig:
    return CadConfig(
        solidworks_template_dir=_solidworks_templates(tmp_path / "modeles"),
        autocad_template_dir=_autocad_templates(tmp_path / "modeles"),
        destination_subfolder=subfolder,
    )


def _project(number: str = "2026-5233", **kwargs: bool) -> ProjectInput:
    return ProjectInput(
        number=parse_project_number(number),
        designation="Escalier helicoidal",
        societe="Client SA",
        **kwargs,  # type: ignore[arg-type]
    )


def _service() -> CadTemplateService:
    return CadTemplateService(
        license_key_loader=lambda: DEMO_LICENSE_KEY,
        manager_factory=open_json_document_manager,
    )


def _statuses(outcome: CadOutcome) -> dict[str, str]:
    return {item.name: item.status for item in outcome.files}


def test_plan_replaces_marker_in_names_and_folders_and_ignores_temporary_files(
    tmp_path: Path,
) -> None:
    template_dir = tmp_path / "modeles"
    (template_dir / "20XX-XXXX-Plans").mkdir(parents=True)
    for name in [
        "20XX-XXXX-PRT-100.SLDPRT",
        "20xx-xxxx-ENS-100.sldasm",
        "~$20XX-XXXX-PRT-100.SLDPRT",
        "20XX-XXXX-ENS-100.bak",
        "Lisez-moi.txt",
        "20XX-XXXX-Plans/20XX-XXXX-ENS-100.SLDDRW",
    ]:
        (template_dir / name).write_bytes(b"x")

    plan = plan_template_copy(template_dir, tmp_path / "projet", "2026-5233")

    assert sorted(
        item.destination.relative_to(tmp_path / "projet").as_posix() for item in plan
    ) == [
        "2026-5233-ENS-100.sldasm",
        "2026-5233-PRT-100.SLDPRT",
        "2026-5233-Plans/2026-5233-ENS-100.SLDDRW",
    ]


def test_copy_new_file_never_overwrites_and_removes_read_only(tmp_path: Path) -> None:
    source = tmp_path / "source.dwg"
    source.write_bytes(b"template")
    source.chmod(stat.S_IREAD)
    existing = tmp_path / "existing.dwg"
    existing.write_bytes(b"project work")

    assert copy_new_file(source, existing) is False
    assert existing.read_bytes() == b"project work"

    copied = tmp_path / "copy" / "copied.dwg"
    assert copy_new_file(source, copied) is True
    assert copied.read_bytes() == b"template"
    assert copied.stat().st_mode & stat.S_IWRITE


def test_copy_failure_reports_onedrive_hint_and_removes_partial_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "20XX-XXXX-ENS-100.dwg"
    source.write_bytes(b"template")
    destination = tmp_path / "2026-5233-ENS-100.dwg"

    def offline_copy(*_args: object) -> None:
        raise OSError("Le fournisseur de fichiers cloud n'est pas en cours d'execution")

    monkeypatch.setattr(templates.shutil, "copyfileobj", offline_copy)

    outcome = CadTemplateService().apply(
        _project(add_autocad=True),
        tmp_path / "projet",
        config=CadConfig(autocad_template_dir=tmp_path),
        initials="LM",
    )

    assert outcome.files[0].status == "error"
    assert "Toujours conserver sur cet appareil" in outcome.files[0].detail
    assert not destination.exists()
    assert not (tmp_path / "projet" / "2026-5233-ENS-100.dwg").exists()


def test_no_option_creates_no_cad_file(tmp_path: Path) -> None:
    project_dir = tmp_path / "projet"
    project_dir.mkdir()

    outcome = _service().apply(_project(), project_dir, config=_config(tmp_path), initials="LM")

    assert outcome == CadOutcome()
    assert list(project_dir.iterdir()) == []


def test_autocad_option_copies_only_the_drawing(tmp_path: Path) -> None:
    project_dir = tmp_path / "projet"

    outcome = _service().apply(
        _project(add_autocad=True),
        project_dir,
        config=_config(tmp_path, subfolder="CAO"),
        initials="LM",
    )

    assert _statuses(outcome) == {"2026-5233-ENS-100.dwg": "created"}
    assert sorted(path.name for path in (project_dir / "CAO").iterdir()) == [
        "2026-5233-ENS-100.dwg",
    ]


def test_solidworks_option_renames_files_rewrites_references_and_properties(
    tmp_path: Path,
) -> None:
    config = _config(tmp_path, subfolder="CAO")
    template_dir = config.solidworks_template_dir
    assert template_dir is not None
    project_dir = tmp_path / "projet"

    outcome = _service().apply(
        _project(add_solidworks=True),
        project_dir,
        config=config,
        initials="LM",
    )

    cad_dir = project_dir / "CAO"
    assert _statuses(outcome) == {
        "2026-5233-ENS-100.SLDASM": "created",
        "2026-5233-ENV-100.SLDPRT": "created",
        "2026-5233-PRT-100.SLDPRT": "created",
        "2026-5233-PRT-200.SLDPRT": "created",
    }
    assert outcome.warnings == ()
    assert outcome.errors == ()
    assembly_properties, references = read_demo_document(cad_dir / "2026-5233-ENS-100.SLDASM")
    assert references == [str(cad_dir / f"2026-5233-{name}") for name in TEMPLATE_PARTS]
    assert assembly_properties == {
        **TEMPLATE_PROPERTIES,
        "Projet": "2026-5233",
        "Client": "Client SA",
        "Auteur": "LM",
        "Description": "Escalier helicoidal",
        "Révision": "A",
        "TAG": "Repère 2026-5233",
    }
    part_properties, _ = read_demo_document(cad_dir / "2026-5233-PRT-100.SLDPRT")
    assert part_properties["Description"] == ""
    assert part_properties["Projet"] == "2026-5233"
    assert part_properties["Fournisseur"] == "Balz Métal SA"
    template_properties, template_references = read_demo_document(
        template_dir / "20XX-XXXX-ENS-100.SLDASM",
    )
    assert template_properties == TEMPLATE_PROPERTIES
    assert template_references == [
        str(template_dir / f"20XX-XXXX-{name}") for name in TEMPLATE_PARTS
    ]


def test_solidworks_references_saved_on_another_computer_are_matched_by_name(
    tmp_path: Path,
) -> None:
    template_dir = _solidworks_templates(
        tmp_path / "modeles",
        reference_root=Path(r"C:\Users\Autre\OneDrive - Balz Metal Sa\11-Racine Solidworks"),
    )
    project_dir = tmp_path / "projet"

    outcome = _service().apply(
        _project(add_solidworks=True),
        project_dir,
        config=CadConfig(solidworks_template_dir=template_dir),
        initials="LM",
    )

    assert outcome.errors == ()
    _, references = read_demo_document(project_dir / "2026-5233-ENS-100.SLDASM")
    assert references == [str(project_dir / f"2026-5233-{name}") for name in TEMPLATE_PARTS]


def test_subproject_files_use_subproject_number_and_folder(tmp_path: Path) -> None:
    project_dir = tmp_path / "2026-5233"
    (project_dir / "2026-5233-2").mkdir(parents=True)

    outcome = _service().apply(
        _project("2026-5233-2", add_solidworks=True, add_autocad=True),
        project_dir,
        config=_config(tmp_path),
        initials="LM",
    )

    target = project_dir / "2026-5233-2"
    assert {item.name for item in outcome.files} == {
        "2026-5233-2-ENS-100.dwg",
        "2026-5233-2-ENS-100.SLDASM",
        "2026-5233-2-ENV-100.SLDPRT",
        "2026-5233-2-PRT-100.SLDPRT",
        "2026-5233-2-PRT-200.SLDPRT",
    }
    properties, references = read_demo_document(target / "2026-5233-2-ENS-100.SLDASM")
    assert properties["Projet"] == "2026-5233-2"
    assert properties["Description"] == "Escalier helicoidal"
    assert references[0] == str(target / "2026-5233-2-ENV-100.SLDPRT")
    assert cad_destination_dir(project_dir, parse_project_number("2026-5233-3"), "CAO") == (
        project_dir / "CAO"
    )


def test_existing_project_files_are_kept_and_missing_ones_added(tmp_path: Path) -> None:
    project_dir = tmp_path / "projet"
    project_dir.mkdir()
    write_demo_document(project_dir / "2026-5233-PRT-100.SLDPRT", {"Projet": "modifie"})
    (project_dir / "2026-5233-ENS-100.dwg").write_bytes(b"plan en cours")

    outcome = _service().apply(
        _project(add_solidworks=True, add_autocad=True),
        project_dir,
        config=_config(tmp_path),
        initials="LM",
    )

    assert _statuses(outcome)["2026-5233-PRT-100.SLDPRT"] == "skipped"
    assert _statuses(outcome)["2026-5233-ENS-100.dwg"] == "skipped"
    assert _statuses(outcome)["2026-5233-PRT-200.SLDPRT"] == "created"
    assert read_demo_document(project_dir / "2026-5233-PRT-100.SLDPRT")[0] == {
        "Projet": "modifie",
    }
    assert (project_dir / "2026-5233-ENS-100.dwg").read_bytes() == b"plan en cours"
    _, references = read_demo_document(project_dir / "2026-5233-ENS-100.SLDASM")
    assert str(project_dir / "2026-5233-PRT-100.SLDPRT") in references


class _StubbornDocument(JsonDocument):
    def replace_reference(self, old_path: str, new_path: str) -> None:
        del old_path, new_path


class _StubbornManager(JsonDocumentManager):
    @contextmanager
    def open_document(self, path: Path, *, read_only: bool = False) -> Iterator[JsonDocument]:
        yield _StubbornDocument(path, read_only=read_only)


def test_reference_left_on_template_is_detected_and_copy_removed(tmp_path: Path) -> None:
    @contextmanager
    def stubborn_factory(_key: str) -> Iterator[_StubbornManager]:
        yield _StubbornManager()

    project_dir = tmp_path / "projet"
    outcome = CadTemplateService(
        license_key_loader=lambda: "key",
        manager_factory=stubborn_factory,
    ).apply(
        _project(add_solidworks=True),
        project_dir,
        config=_config(tmp_path),
        initials="LM",
    )

    assembly = next(item for item in outcome.files if item.name.endswith(".SLDASM"))
    assert assembly.status == "error"
    assert "encore liee(s) au dossier modele" in assembly.detail
    assert "20XX-XXXX-ENV-100.SLDPRT" in assembly.detail
    assert not (project_dir / "2026-5233-ENS-100.SLDASM").exists()
    assert (project_dir / "2026-5233-PRT-100.SLDPRT").exists()


@pytest.mark.parametrize(
    "factory",
    [
        pytest.param(open_json_document_manager, id="invalid-key"),
        pytest.param(open_document_manager, id="document-manager-absent"),
    ],
)
def test_without_document_manager_parts_are_copied_but_not_assemblies(
    tmp_path: Path,
    factory: object,
) -> None:
    if factory is open_document_manager and sys.platform == "win32":
        pytest.skip("Document Manager peut etre installe sur ce poste")
    project_dir = tmp_path / "projet"

    outcome = CadTemplateService(
        license_key_loader=lambda: "cle-invalide",
        manager_factory=factory,  # type: ignore[arg-type]
    ).apply(
        _project(add_solidworks=True, add_autocad=True),
        project_dir,
        config=_config(tmp_path),
        initials="LM",
    )

    assert _statuses(outcome) == {
        "2026-5233-ENS-100.dwg": "created",
        "2026-5233-ENS-100.SLDASM": "skipped",
        "2026-5233-ENV-100.SLDPRT": "created",
        "2026-5233-PRT-100.SLDPRT": "created",
        "2026-5233-PRT-200.SLDPRT": "created",
    }
    assert not (project_dir / "2026-5233-ENS-100.SLDASM").exists()
    assert read_demo_document(project_dir / "2026-5233-PRT-100.SLDPRT")[0] == TEMPLATE_PROPERTIES
    assert len(outcome.warnings) == 1
    assert "Document Manager indisponible" in outcome.warnings[0]
    assert "cle-invalide" not in outcome.warnings[0]


def test_missing_license_key_is_reported_as_unavailable() -> None:
    with pytest.raises(CadUnavailableError), open_document_manager("  "):
        pass


def test_missing_template_folder_is_an_error_for_that_option_only(tmp_path: Path) -> None:
    outcome = _service().apply(
        _project(add_solidworks=True, add_autocad=True),
        tmp_path / "projet",
        config=CadConfig(autocad_template_dir=_autocad_templates(tmp_path)),
        initials="LM",
    )

    assert _statuses(outcome) == {"2026-5233-ENS-100.dwg": "created"}
    assert len(outcome.errors) == 1
    assert outcome.errors[0].startswith("SolidWorks : Dossier modele SolidWorks non configure")


def test_property_updates_follow_mapping() -> None:
    names = CadPropertyNames(client="Maitre d'ouvrage", revision="Rev", revision_defaut="0")
    existing = {
        "Projet": "",
        "Maitre d'ouvrage": "",
        "Client": "Ancien",
        "Rev": "",
        "Auteur": "XX",
        "Note": "voir 20xx-xxxx",
        "Fournisseur": "Balz Métal SA",
    }

    updates = compute_property_updates(
        existing,
        number="2026-5233",
        societe="Client SA",
        initials="LM",
        designation="Garde-corps",
        write_description=False,
        names=names,
    )

    assert updates == {
        "Projet": "2026-5233",
        "Maitre d'ouvrage": "Client SA",
        "Rev": "0",
        "Auteur": "LM",
        "Note": "voir 2026-5233",
    }


def test_property_updates_keep_existing_revision_and_description_scope() -> None:
    updates = compute_property_updates(
        {"Projet": "Existant", "Révision": "C", "Description": "modele"},
        number="2026-5233",
        societe="",
        initials="",
        designation="Escalier",
        write_description=True,
        names=CadPropertyNames(),
    )

    assert updates == {"Description": "Escalier"}


def test_template_availability_explains_why_an_option_is_disabled(tmp_path: Path) -> None:
    missing = template_availability(tmp_path / "absent", "AutoCAD")
    assert not missing.available
    assert "introuvable" in missing.reason
    assert "non configure" in template_availability(None, "SolidWorks").reason
    assert template_availability(tmp_path, "SolidWorks").available


@pytest.mark.parametrize("value", ["../autre", "C:/CAO", "/CAO"])
def test_cad_subfolder_must_stay_inside_project(value: str) -> None:
    with pytest.raises(ValidationError):
        CadConfig(destination_subfolder=value)


def test_cad_subfolder_is_normalized() -> None:
    assert CadConfig(destination_subfolder=" 03-CAO\\SolidWorks/ ").destination_subfolder == (
        "03-CAO/SolidWorks"
    )


def test_license_key_uses_credential_store_only() -> None:
    storage = SolidWorksLicenseStorage()

    assert storage.load() == ""
    assert storage.save(" Balz:swdocmgr_general-00000 ")
    assert storage.load() == "Balz:swdocmgr_general-00000"
    assert storage.clear()
    assert storage.load() == ""


class _FailingCad:
    def apply(self, *_args: object, **_kwargs: object) -> CadOutcome:
        raise OSError("disque plein")


class _RecordingRepertoire:
    async def upsert_project(self, project: ProjectInput, *, force_overwrite: bool = False) -> None:
        del project, force_overwrite


def _project_service(tmp_path: Path, cad: object) -> ProjectService:
    config = AppConfig()
    config.user.initials = "LM"
    config.paths.racine_projets = tmp_path / "clients"
    config.paths.dossier_reference = tmp_path / "reference"
    config.paths.dossier_reference.mkdir()
    workbook = Workbook()
    workbook.save(config.paths.dossier_reference / "modele fiche.xlsx")
    workbook.close()
    config.cad = _config(tmp_path, subfolder="CAO")
    return ProjectService(
        config=config,
        fiche_service=FicheService(),
        repertoire_service=_RecordingRepertoire(),  # type: ignore[arg-type]
        cad=cad,  # type: ignore[arg-type]
    )


@pytest.mark.asyncio
async def test_project_creation_reports_cad_files(tmp_path: Path) -> None:
    service = _project_service(tmp_path, _service())

    result = await service.create_project(_project(add_solidworks=True, add_autocad=True))

    cad_dir = Path(result.project_dir) / "CAO"
    assert result.project_dir_created
    assert len([item for item in result.cad_files if item.status == "created"]) == 5
    assert result.cad_error is None
    assert (cad_dir / "2026-5233-ENS-100.SLDASM").exists()


@pytest.mark.asyncio
async def test_project_creation_without_cad_option_adds_no_cad_file(tmp_path: Path) -> None:
    service = _project_service(tmp_path, _service())

    result = await service.create_project(_project())

    assert result.cad_files == ()
    assert not (Path(result.project_dir) / "CAO").exists()


@pytest.mark.asyncio
async def test_cad_failure_does_not_cancel_project_creation(tmp_path: Path) -> None:
    service = _project_service(tmp_path, _FailingCad())

    result = await service.create_project(_project(add_autocad=True))

    assert result.project_dir_created
    assert result.fiche_path is not None
    assert result.cad_error == "disque plein"


@pytest.mark.asyncio
async def test_update_existing_project_adds_only_missing_cad_files(tmp_path: Path) -> None:
    service = _project_service(tmp_path, _service())
    created = await service.create_project(_project(add_autocad=True))
    drawing = Path(created.project_dir) / "CAO" / "2026-5233-ENS-100.dwg"
    drawing.write_bytes(b"plan modifie")

    result = await service.update_project(_project(add_solidworks=True, add_autocad=True))

    assert _statuses(CadOutcome(files=result.cad_files))["2026-5233-ENS-100.dwg"] == "skipped"
    assert drawing.read_bytes() == b"plan modifie"
    assert (Path(created.project_dir) / "CAO" / "2026-5233-ENS-100.SLDASM").exists()


@pytest.mark.asyncio
async def test_subproject_creation_adds_cad_files(tmp_path: Path) -> None:
    service = _project_service(tmp_path, _service())
    await service.create_project(_project())

    result = await service.create_subproject(_project("2026-5233-2", add_autocad=True))

    assert [item.name for item in result.cad_files] == ["2026-5233-2-ENS-100.dwg"]

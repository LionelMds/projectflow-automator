from __future__ import annotations

from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from projectflow.core.fiche_service import FicheService
from projectflow.core.numero import parse_project_number
from projectflow.core.sortie_service import SortieDossierService, SortieSelection
from projectflow.exceptions import ProjectCreationError


def test_discover_collects_fiche_measurements_photos_and_plans(tmp_path: Path) -> None:
    project_dir = tmp_path / "2026" / "2026-5093"
    project_dir.mkdir(parents=True)
    (project_dir / "2026-5093 - Fiche dossier clients.xlsx").touch()
    (project_dir / "Prise de cote initiale.pdf").touch()
    photos = project_dir / "photos"
    photos.mkdir()
    (photos / "photo-01.jpg").touch()
    plans = project_dir / "Plans" / "Plan d'execution"
    plans.mkdir(parents=True)
    (plans / "Plan facade.pdf").touch()

    inventory = SortieDossierService(FicheService()).discover(
        project_dir,
        parse_project_number("2026-5093"),
    )

    assert [candidate.path.name for candidate in inventory.fiches] == [
        "2026-5093 - Fiche dossier clients.xlsx",
    ]
    assert [candidate.path.name for candidate in inventory.mesure_pdfs] == [
        "Prise de cote initiale.pdf",
    ]
    assert [candidate.path.name for candidate in inventory.photos] == ["photo-01.jpg"]
    assert [candidate.path.name for candidate in inventory.plans] == ["Plan facade.pdf"]


def test_discover_also_finds_numbered_nested_folder(tmp_path: Path) -> None:
    project_dir = tmp_path / "2026" / "2026-5093"
    nested = project_dir / "2026-5093"
    nested.mkdir(parents=True)
    workbook = Workbook()
    workbook.save(nested / "2026-5093 - Fiche dossier clients.xlsx")
    workbook.close()
    (nested / "Prise de cote.pdf").touch()

    inventory = SortieDossierService(FicheService()).discover(
        project_dir,
        parse_project_number("2026-5093"),
    )

    assert len(inventory.fiches) == 1
    assert len(inventory.mesure_pdfs) == 1


def test_selection_rejects_missing_files(tmp_path: Path) -> None:
    selection = SortieSelection(fiche_path=tmp_path / "missing.xlsx")

    with pytest.raises(ProjectCreationError, match="introuvable"):
        selection.validate()

    workbook = Workbook()
    fiche = tmp_path / "fiche.xlsx"
    workbook.save(fiche)
    workbook.close()
    SortieSelection(fiche_path=fiche).validate()


@pytest.mark.parametrize("e2", ["fiche d'atelier le", "fiche d'atelier le 01.02.2020"])
def test_create_output_folder_copies_documents_without_modifying_sources(
    tmp_path: Path,
    e2: str,
) -> None:
    project_dir = tmp_path / "2026" / "2026-5093"
    project_dir.mkdir(parents=True)
    fiche = project_dir / "fiche.xlsx"
    mesure = project_dir / "cote.pdf"
    photo = project_dir / "photo.jpg"
    photo_without_rotation = project_dir / "photo-2.jpg"
    plan = project_dir / "plan.pdf"
    for path in (mesure, photo, photo_without_rotation, plan):
        path.write_bytes(path.name.encode())
    workbook = Workbook()
    workbook.active["E2"] = e2
    workbook.active["B9"] = date(2024, 3, 4)
    workbook.active["C6"] = "Alice Martin"
    workbook.active["C9"] = "LM"
    workbook.save(fiche)
    workbook.close()
    original_photo = photo.read_bytes()
    original_fiche = fiche.read_bytes()

    selection = SortieSelection(
        fiche_path=fiche,
        mesure_pdf_path=mesure,
        photo_paths=(photo, photo_without_rotation),
        plan_paths=(plan,),
    )
    service = SortieDossierService(
        FicheService(today=lambda: date(2026, 7, 15)),
        now=lambda: datetime(2026, 7, 15, 10, 11, 12, tzinfo=UTC),
    )

    output_dir = service.create_output_folder(
        project_dir,
        parse_project_number("2026-5093"),
        selection,
    )

    assert output_dir.name == "2026-5093 - Sortie dossier - 20260715-101112"
    copied_fiche = output_dir / "01 - Fiche dossier" / fiche.name
    assert copied_fiche.exists()
    output_workbook = load_workbook(copied_fiche)
    assert output_workbook.active["E2"].value == "fiche d'atelier le 15.07.2026"
    assert output_workbook.active["B9"].value.date() == date(2024, 3, 4)
    assert output_workbook.active["C6"].value == "Alice Martin"
    assert output_workbook.active["C9"].value == "LM"
    output_workbook.close()
    source_workbook = load_workbook(fiche)
    assert source_workbook.active["E2"].value == e2
    assert source_workbook.active["B9"].value.date() == date(2024, 3, 4)
    assert source_workbook.active["C6"].value == "Alice Martin"
    assert source_workbook.active["C9"].value == "LM"
    source_workbook.close()
    assert fiche.read_bytes() == original_fiche
    assert (output_dir / "02 - Prise de cote" / mesure.name).exists()
    copied_photo = output_dir / "03 - Photos" / photo.name
    assert copied_photo.read_bytes() == original_photo
    assert (
        output_dir / "03 - Photos" / photo_without_rotation.name
    ).read_bytes() == photo_without_rotation.name.encode()
    assert (output_dir / "04 - Plans" / plan.name).exists()
    assert photo.read_bytes() == original_photo


def test_create_output_folder_uses_a_suffix_when_timestamp_folder_exists(tmp_path: Path) -> None:
    project_dir = tmp_path / "2026" / "2026-5093"
    project_dir.mkdir(parents=True)
    fiche = project_dir / "fiche.xlsx"
    workbook = Workbook()
    workbook.save(fiche)
    workbook.close()
    base = project_dir / "Sorties dossier" / "2026-5093 - Sortie dossier - 20260715-101112"
    base.mkdir(parents=True)

    output_dir = SortieDossierService(
        FicheService(),
        now=lambda: datetime(2026, 7, 15, 10, 11, 12, tzinfo=UTC),
    ).create_output_folder(
        project_dir,
        parse_project_number("2026-5093"),
        SortieSelection(fiche_path=fiche),
    )

    assert output_dir.name.endswith("(2)")


def test_discover_finds_accented_plan_folder_and_reports_directories(tmp_path: Path) -> None:
    project_dir = tmp_path / "2026" / "2026-5093"
    plans = project_dir / "Plans" / "Plan d'exécution"
    plans.mkdir(parents=True)
    (plans / "2026-5093-ENS-100.pdf").touch()
    photos = project_dir / "Photos"
    photos.mkdir()
    (project_dir / "2026-5093 - Fiche dossier clients.xlsx").touch()

    inventory = SortieDossierService(FicheService()).discover(
        project_dir,
        parse_project_number("2026-5093"),
    )

    assert [candidate.path.name for candidate in inventory.plans] == ["2026-5093-ENS-100.pdf"]
    assert inventory.mesure_pdfs == ()
    assert inventory.plan_directory == plans.resolve()
    assert inventory.photo_directory == photos.resolve()
